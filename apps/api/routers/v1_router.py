"""
Clisonix V1 API — Command + Report + Reader pipeline
Spec: CLISONIX_SELLABLE_V1_SPEC.md
"""

import io
import json
import os
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, Tuple, cast

from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    Header,
    HTTPException,
    Request,
    UploadFile,
)
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel

# ---------------------------------------------------------------------------
# API Key enforcement (imports from api_monetization in parent package)
# ---------------------------------------------------------------------------
try:
    sys.path.insert(0, str(Path(__file__).parent.parent))
    from api_monetization import (
        PLAN_LIMITS,
        APIKey,
        APIUsage,
        SessionLocal,
        hash_api_key,
        validate_api_key,
    )
    _MONETIZATION_AVAILABLE = True
except ImportError:
    PLAN_LIMITS = {}
    APIKey = None
    APIUsage = None
    SessionLocal = None
    hash_api_key = None
    validate_api_key = None
    _MONETIZATION_AVAILABLE = False


def _record_usage(api_key_obj: Any, endpoint: str) -> None:
    """Increment usage counter for a validated key."""
    if not _MONETIZATION_AVAILABLE or SessionLocal is None or APIUsage is None:
        return
    db = None
    try:
        db = SessionLocal()
        usage_model = APIUsage
        today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        usage = db.query(usage_model).filter(
            usage_model.api_key_id == api_key_obj.id,
            usage_model.date == today,
        ).first()
        if usage:
            usage = cast(Any, usage)
            usage.requests += 1
        else:
            usage = usage_model(
                id=str(uuid.uuid4()),
                api_key_id=api_key_obj.id,
                user_id=api_key_obj.user_id,
                date=today,
                requests=1,
                endpoint=endpoint,
            )
            db.add(usage)
        api_key_obj.last_used = datetime.now(timezone.utc)
        db.commit()
    except Exception:
        pass
    finally:
        try:
            if db is not None:
                db.close()
        except Exception:
            pass


def require_api_key(x_api_key: Optional[str] = Header(default=None, alias="X-API-Key")) -> Optional[Any]:
    """
    FastAPI dependency — enforces API key on V1 endpoints.
    - If monetization module not available: passes through (dev mode).
    - If no key provided: 401.
    - If invalid / rate-limited: 401 / 429.
    """
    if not _MONETIZATION_AVAILABLE or SessionLocal is None or validate_api_key is None:
        return None  # dev mode: allow all
    if not x_api_key:
        raise HTTPException(
            status_code=401,
            detail="Missing X-API-Key header. Get a key at /api/v1/api-access/keys/create",
        )
    db = SessionLocal()
    try:
        is_valid, key_obj, error = validate_api_key(x_api_key, db)
        if not is_valid:
            if key_obj is not None:
                # Key exists but rate-limited
                raise HTTPException(status_code=429, detail=error)
            raise HTTPException(status_code=401, detail=error)
        # Detach so we can use key_obj after session closes
        db.expunge(key_obj)
        return key_obj
    finally:
        db.close()

# ---------------------------------------------------------------------------
# In-memory job store (replace with DB/Redis in production)
# ---------------------------------------------------------------------------
_JOBS: Dict[str, Dict[str, Any]] = {}

OUTPUT_DIR = Path(os.environ.get("REPORT_OUTPUT_DIR", "/tmp/clisonix_reports"))
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

router = APIRouter(prefix="/api/v1", tags=["V1"])

# ---------------------------------------------------------------------------
# Shared models
# ---------------------------------------------------------------------------

class CommandRequest(BaseModel):
    command: str
    output_format: Literal["excel", "word", "pdf", "pptx"] = "pdf"
    template: Optional[str] = None  # "sales_report_pack" | "ops_weekly_brief" | "research_summary_pack"
    sources: Optional[List[str]] = None  # source_ids from /reader/* endpoints
    context: Optional[Dict[str, Any]] = None


class ReportJobRequest(BaseModel):
    command: str
    output_format: Literal["excel", "word", "pdf", "pptx"] = "pdf"
    template: Optional[str] = None
    sources: Optional[List[str]] = None  # URLs or file refs already ingested


class IngestUrlRequest(BaseModel):
    url: str
    label: Optional[str] = None


# ---------------------------------------------------------------------------
# Job helpers
# ---------------------------------------------------------------------------

def _new_job(command: str, output_format: str, template: Optional[str]) -> Dict[str, Any]:
    job_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    job = {
        "job_id": job_id,
        "status": "queued",
        "command": command,
        "output_format": output_format,
        "template": template,
        "created_at": now,
        "updated_at": now,
        "output_path": None,
        "error": None,
    }
    _JOBS[job_id] = job
    return job


def _update_job(job_id: str, **kwargs: Any) -> None:
    if job_id in _JOBS:
        _JOBS[job_id].update(kwargs)
        _JOBS[job_id]["updated_at"] = datetime.now(timezone.utc).isoformat()


def _job_or_404(job_id: str) -> Dict[str, Any]:
    job = _JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail=f"Job {job_id} not found")
    return job


def _resolve_source_previews(source_ids: Optional[List[str]]) -> List[str]:
    """Resolve source previews for report content; return empty when no resolvable sources."""
    previews: List[str] = []
    ids = source_ids or []
    for source_id in ids:
        source = _SOURCES.get(source_id)
        if not source:
            continue
        preview = str(source.get("preview", "")).strip()
        if preview:
            previews.append(preview)
    return previews


def _extract_world_bank_preview(raw: str) -> Optional[str]:
    """Parse World Bank payload shape [metadata, data] and return a factual preview."""
    try:
        payload = json.loads(raw)
    except Exception:
        return None

    if not isinstance(payload, list) or len(payload) < 2:
        return None

    metadata = payload[0]
    rows = payload[1]
    if not isinstance(metadata, dict) or not isinstance(rows, list):
        return None

    latest_non_null: Optional[Dict[str, Any]] = None
    for row in rows:
        if isinstance(row, dict) and row.get("value") is not None:
            latest_non_null = row
            break

    if latest_non_null is None:
        return None

    indicator_id = ""
    indicator_name = ""
    country_name = ""
    if isinstance(latest_non_null.get("indicator"), dict):
        indicator_id = str(latest_non_null["indicator"].get("id") or "").strip()
        indicator_name = str(latest_non_null["indicator"].get("value") or "").strip()
    if isinstance(latest_non_null.get("country"), dict):
        country_name = str(latest_non_null["country"].get("value") or "").strip()

    top_year = str(rows[0].get("date")) if rows and isinstance(rows[0], dict) else ""
    top_value = rows[0].get("value") if rows and isinstance(rows[0], dict) else None
    latest_year = str(latest_non_null.get("date") or "")

    latest_value_raw = latest_non_null.get("value")
    if latest_value_raw is None:
        latest_value_str = "not available"
    else:
        try:
            latest_value = float(latest_value_raw)
            latest_value_str = f"{latest_value:.4f}"
        except Exception:
            latest_value_str = str(latest_value_raw)

    source_id = str(metadata.get("sourceid") or "").strip()
    last_updated = str(metadata.get("lastupdated") or "").strip()

    recent_lines: List[str] = []
    for row in rows[:6]:
        if not isinstance(row, dict):
            continue
        year = str(row.get("date") or "")
        value = row.get("value")
        if value is None:
            recent_lines.append(f"{year}: not available")
            continue
        try:
            recent_lines.append(f"{year}: {float(value):.4f}")
        except Exception:
            recent_lines.append(f"{year}: {value}")

    header_parts = ["World Bank data detected"]
    if indicator_id:
        header_parts.append(f"indicator={indicator_id}")
    if indicator_name:
        header_parts.append(indicator_name)
    if country_name:
        header_parts.append(f"scope={country_name}")

    lines = [" | ".join(header_parts)]
    lines.append(
        f"Latest available value: {latest_year} = {latest_value_str}"
        + (f" (top year {top_year} unavailable)" if top_year and top_value is None else "")
    )
    if source_id or last_updated:
        lines.append(f"Source: WDI/{source_id or '-'} | Last updated: {last_updated or '-'}")
    if recent_lines:
        lines.append("Recent series: " + "; ".join(recent_lines))

    return " ".join(lines)[:1500]


def _build_source_preview(raw: str) -> str:
    world_bank_preview = _extract_world_bank_preview(raw)
    if world_bank_preview:
        return world_bank_preview

    # Generic fallback: strip HTML tags and compress whitespace.
    import re

    clean = re.sub(r"<[^>]+>", " ", raw)
    clean = re.sub(r"\s+", " ", clean).strip()
    return clean[:500]


# ---------------------------------------------------------------------------
# Background generator
# ---------------------------------------------------------------------------

def _generate_report(
    job_id: str,
    output_format: str,
    command: str,
    template: Optional[str],
    source_previews: List[str],
) -> None:
    """Runs in background. Generates a real (non-fake) report file."""
    _update_job(job_id, status="processing")
    try:
        out_file = OUTPUT_DIR / f"{job_id}.{output_format}"

        if output_format == "excel":
            _gen_excel(out_file, command, template, source_previews)
        elif output_format == "word":
            _gen_word(out_file, command, template, source_previews)
        elif output_format == "pdf":
            _gen_pdf(out_file, command, template, source_previews)
        elif output_format == "pptx":
            _gen_pptx(out_file, command, template, source_previews)
        else:
            raise ValueError(f"Unsupported format: {output_format}")

        _update_job(job_id, status="complete", output_path=str(out_file))
    except Exception as exc:  # broad catch — surface error to client, never fake data
        _update_job(job_id, status="failed", error=str(exc))


def _gen_excel(path: Path, command: str, template: Optional[str], source_previews: List[str]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise RuntimeError("openpyxl is required for Excel output. Install it: pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("Workbook active worksheet not available")
    ws.title = "Report"
    ws.append(["Clisonix V1 Report"])
    ws.append(["Generated", datetime.now(timezone.utc).isoformat()])
    ws.append(["Command", command])
    ws.append(["Template", template or "none"])
    ws.append([])
    ws.append(["Section", "Value"])
    ws.append(["Status", "Report generation complete using ingested real sources"])
    ws.append([])
    ws.append(["Source", "Preview Snippet"])
    for i, preview in enumerate(source_previews[:3], start=1):
        ws.append([f"Source {i}", preview[:180]])
    wb.save(path)


def _gen_word(path: Path, command: str, template: Optional[str], source_previews: List[str]) -> None:
    try:
        from docx import Document  # type: ignore[import-not-found]
    except ImportError as e:
        raise RuntimeError("python-docx is required for Word output. Install it: pip install python-docx") from e

    doc = Document()
    doc.add_heading("Clisonix V1 Report", 0)
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).isoformat()}")
    doc.add_paragraph(f"Command: {command}")
    doc.add_paragraph(f"Template: {template or 'none'}")
    doc.add_paragraph("This report is generated from ingested real sources.")
    for i, preview in enumerate(source_previews[:3], start=1):
        doc.add_paragraph(f"Source {i}: {preview[:300]}")
    doc.save(path)


def _gen_pdf(path: Path, command: str, template: Optional[str], source_previews: List[str]) -> None:
    try:
        from reportlab.lib.pagesizes import A4  # type: ignore[import-not-found]
        from reportlab.pdfgen import (
            canvas as rl_canvas,  # type: ignore[import-not-found]  # pyright: ignore[reportMissingModuleSource]
        )
    except ImportError:
        # Fallback: minimal raw PDF (no fake data — just metadata)
        _write_minimal_pdf(path, command, template, source_previews)
        return

    c = rl_canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, 800, "Clisonix V1 Report")
    c.setFont("Helvetica", 11)
    c.drawString(50, 775, f"Generated: {datetime.now(timezone.utc).isoformat()}")
    c.drawString(50, 755, f"Command: {command}")
    c.drawString(50, 735, f"Template: {template or 'none'}")
    c.drawString(50, 700, "Generated from ingested real sources")
    y = 675
    for i, preview in enumerate(source_previews[:2], start=1):
        c.drawString(50, y, f"Source {i}: {preview[:95]}")
        y -= 18
    c.save()


def _write_minimal_pdf(path: Path, command: str, template: Optional[str], source_previews: List[str]) -> None:
    """Minimal valid PDF without external libraries, with visible text content."""

    def _pdf_escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    ts = datetime.now(timezone.utc).isoformat()
    text_lines = [
        "Clisonix V1 Report",
        f"Generated: {ts}",
        f"Command: {command[:120]}",
        f"Template: {template or 'none'}",
        "Generated from ingested real sources",
    ]
    for i, preview in enumerate(source_previews[:2], start=1):
        text_lines.append(f"Source {i}: {preview[:140]}")

    stream_lines = [
        "BT",
        "/F1 12 Tf",
        "50 800 Td",
        "16 TL",
    ]
    for idx, line in enumerate(text_lines):
        escaped = _pdf_escape(line)
        if idx == 0:
            stream_lines.append(f"({escaped}) Tj")
        else:
            stream_lines.append("T*")
            stream_lines.append(f"({escaped}) Tj")
    stream_lines.append("ET")

    stream_bytes = ("\n".join(stream_lines) + "\n").encode("latin-1", errors="replace")

    obj1 = b"<< /Type /Catalog /Pages 2 0 R >>"
    obj2 = b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>"
    obj3 = (
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
    )
    obj4 = (
        f"<< /Length {len(stream_bytes)} >>\nstream\n".encode("ascii")
        + stream_bytes
        + b"endstream"
    )
    obj5 = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    objects = [obj1, obj2, obj3, obj4, obj5]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out.extend(f"{idx} 0 obj\n".encode("ascii"))
        out.extend(obj)
        out.extend(b"\nendobj\n")

    xref_offset = len(out)
    out.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    out.extend(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        out.extend(f"{off:010d} 00000 n \n".encode("ascii"))
    out.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode("ascii"))
    out.extend(f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii"))

    path.write_bytes(bytes(out))


def _gen_pptx(path: Path, command: str, template: Optional[str], source_previews: List[str]) -> None:
    try:
        from pptx import Presentation  # type: ignore[import-not-found]
        from pptx.util import Inches, Pt  # type: ignore[import-not-found]
    except ImportError as e:
        raise RuntimeError("python-pptx is required for PPT output. Install it: pip install python-pptx") from e

    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Clisonix V1 Report"
    slide.placeholders[1].text = (
        f"Command: {command}\nTemplate: {template or 'none'}\n"
        f"Generated: {datetime.now(timezone.utc).isoformat()}"
    )

    details_layout = prs.slide_layouts[1]
    details = prs.slides.add_slide(details_layout)
    details.shapes.title.text = "Real Source Highlights"
    details.placeholders[1].text = "\n".join(
        [f"Source {i}: {preview[:180]}" for i, preview in enumerate(source_previews[:3], start=1)]
    ) or "No source snippets available"
    prs.save(path)


# ---------------------------------------------------------------------------
# In-memory reader sources store
# ---------------------------------------------------------------------------
_SOURCES: Dict[str, Dict[str, Any]] = {}


# ===========================================================================
# ROUTES — Commands
# ===========================================================================

@router.post("/commands/execute")
async def execute_command(
    body: CommandRequest,
    background_tasks: BackgroundTasks,
    api_key_obj: Any = Depends(require_api_key),
) -> JSONResponse:
    """
    Accept a natural-language command and immediately create + queue a report job.
    Returns job metadata so client can poll /reports/jobs/{job_id}.
    """
    source_ids = body.sources or []
    source_previews = _resolve_source_previews(source_ids)
    if not source_previews:
        raise HTTPException(
            status_code=422,
            detail="No real data sources provided. Ingest URL/file first and include source_ids.",
        )
    if api_key_obj is not None:
        _record_usage(api_key_obj, "/api/v1/commands/execute")
    job = _new_job(body.command, body.output_format, body.template)
    background_tasks.add_task(
        _generate_report,
        job["job_id"],
        body.output_format,
        body.command,
        body.template,
        source_previews,
    )
    return JSONResponse(status_code=202, content={"job": job})


# ===========================================================================
# ROUTES — Report Jobs
# ===========================================================================

@router.post("/reports/jobs")
async def create_report_job(
    body: ReportJobRequest,
    background_tasks: BackgroundTasks,
    api_key_obj: Any = Depends(require_api_key),
) -> JSONResponse:
    source_ids = body.sources or []
    source_previews = _resolve_source_previews(source_ids)
    if not source_previews:
        raise HTTPException(
            status_code=422,
            detail="No real data sources provided. Ingest URL/file first and include source_ids.",
        )
    if api_key_obj is not None:
        _record_usage(api_key_obj, "/api/v1/reports/jobs")
    job = _new_job(body.command, body.output_format, body.template)
    background_tasks.add_task(
        _generate_report,
        job["job_id"],
        body.output_format,
        body.command,
        body.template,
        source_previews,
    )
    return JSONResponse(status_code=202, content={"job": job})


@router.get("/reports/jobs/{job_id}")
async def get_report_job(
    job_id: str,
    api_key_obj: Any = Depends(require_api_key),
) -> JSONResponse:
    if api_key_obj is not None:
        _record_usage(api_key_obj, f"/api/v1/reports/jobs/{job_id}")
    job = _job_or_404(job_id)
    safe = {k: v for k, v in job.items() if k != "output_path"}
    safe["download_url"] = f"/api/v1/reports/jobs/{job_id}/download" if job["status"] == "complete" else None
    return JSONResponse(content={"job": safe})


@router.get("/reports/jobs/{job_id}/download")
async def download_report(
    job_id: str,
    api_key_obj: Any = Depends(require_api_key),
) -> FileResponse:
    job = _job_or_404(job_id)
    if job["status"] != "complete":
        raise HTTPException(status_code=409, detail=f"Job status is '{job['status']}', not ready for download")
    path = Path(job["output_path"])
    if not path.exists():
        raise HTTPException(status_code=404, detail="Output file not found on server")
    media_map = {
        "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pdf": "application/pdf",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    fmt = job["output_format"]
    return FileResponse(
        path=str(path),
        media_type=media_map.get(fmt, "application/octet-stream"),
        filename=f"clisonix_report_{job_id}.{fmt}",
    )


# ===========================================================================
# ROUTES — Output-specific (convenience shortcuts)
# ===========================================================================

class OutputRequest(BaseModel):
    command: str
    template: Optional[str] = None
    sources: Optional[List[str]] = None


def _output_shortcut(fmt: str):
    async def handler(
        body: OutputRequest,
        background_tasks: BackgroundTasks,
        api_key_obj: Any = Depends(require_api_key),
    ) -> JSONResponse:
        source_ids = body.sources or []
        source_previews = _resolve_source_previews(source_ids)
        if not source_previews:
            raise HTTPException(
                status_code=422,
                detail="No real data sources provided. Ingest URL/file first and include source_ids.",
            )
        if api_key_obj is not None:
            _record_usage(api_key_obj, f"/api/v1/reports/{fmt}")
        job = _new_job(body.command, fmt, body.template)
        background_tasks.add_task(_generate_report, job["job_id"], fmt, body.command, body.template, source_previews)
        return JSONResponse(status_code=202, content={"job": job})
    handler.__name__ = f"create_{fmt}_report"
    return handler


router.post("/reports/excel")(_output_shortcut("excel"))
router.post("/reports/word")(_output_shortcut("word"))
router.post("/reports/pdf")(_output_shortcut("pdf"))
router.post("/reports/pptx")(_output_shortcut("pptx"))


# ===========================================================================
# ROUTES — Reader / Data Ingestion
# ===========================================================================

@router.post("/reader/ingest-url")
async def ingest_url(body: IngestUrlRequest) -> JSONResponse:
    """Fetch and store summary of a URL for use in report generation."""
    import urllib.error
    import urllib.request

    source_id = str(uuid.uuid4())
    label = body.label or body.url[:60]

    try:
        req = urllib.request.Request(
            body.url,
            headers={"User-Agent": "ClisonixReader/1.0"},
        )
        with urllib.request.urlopen(req, timeout=10) as resp:
            raw = resp.read(32768).decode("utf-8", errors="replace")
        preview = _build_source_preview(raw)
    except urllib.error.URLError as exc:
        raise HTTPException(status_code=502, detail=f"Failed to fetch URL: {exc.reason}")
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Failed to fetch URL: {exc}")

    _SOURCES[source_id] = {
        "source_id": source_id,
        "type": "url",
        "url": body.url,
        "label": label,
        "preview": preview,
        "ingested_at": datetime.now(timezone.utc).isoformat(),
    }
    return JSONResponse(status_code=201, content={"source": _SOURCES[source_id]})


@router.post("/reader/ingest-file")
async def ingest_file(file: UploadFile = File(...)) -> JSONResponse:
    """Accept an uploaded file and store reference for report generation."""
    source_id = str(uuid.uuid4())
    content = await file.read()
    preview_bytes = content[:500]

    # Try to decode as text for preview
    try:
        preview = preview_bytes.decode("utf-8", errors="replace")
    except Exception:
        preview = f"[binary file, {len(content)} bytes]"

    out_path = OUTPUT_DIR / f"src_{source_id}_{file.filename}"
    out_path.write_bytes(content)

    _SOURCES[source_id] = {
        "source_id": source_id,
        "type": "file",
        "filename": file.filename,
        "size_bytes": len(content),
        "preview": preview[:300],
        "stored_path": str(out_path),
        "ingested_at": datetime.now(timezone.utc).isoformat(),
    }
    return JSONResponse(status_code=201, content={"source": {k: v for k, v in _SOURCES[source_id].items() if k != "stored_path"}})


@router.get("/reader/sources")
async def list_sources() -> JSONResponse:
    safe = [{k: v for k, v in s.items() if k != "stored_path"} for s in _SOURCES.values()]
    return JSONResponse(content={"sources": safe, "count": len(safe)})
