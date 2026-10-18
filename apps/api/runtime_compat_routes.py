"""Compatibility routes backed by configured services and existing handlers."""
from __future__ import annotations

import asyncio
import os
import io
import json
import textwrap
from datetime import datetime, timezone

import httpx
from fastapi import APIRouter, FastAPI, HTTPException, Query, Request
from fastapi.responses import JSONResponse, Response

router = APIRouter(tags=["runtime-compatibility"])
SERVICE_ENV = {
    # ALBA has no default: older compose files point it at a stub.
    "alba": ("ALBA_SERVICE_URL", "ALBA_URL"),
    "albi": ("ALBI_SERVICE_URL", "ALBI_URL"),
    "jona": ("JONA_SERVICE_URL", "JONA_URL"),
    "ocean": ("OCEAN_INTERNAL_URL", "OCEAN_CORE_URL"),
    "reporting": ("REPORTING_INTERNAL_URL", "REPORTING_SERVICE_URL"),
    "marketplace": ("MARKETPLACE_INTERNAL_URL", "MARKETPLACE_URL"),
}
DEFAULTS = {
    "albi": "http://clisonix-albi:6680",
    "jona": "http://clisonix-jona:7777",
    "ocean": "http://clisonix-ocean-core:8030",
    "reporting": "http://clisonix-reporting:8001",
}


def service_base(service: str) -> str:
    if service not in SERVICE_ENV:
        raise HTTPException(404, "Unknown service")
    base = next((os.environ[key].strip() for key in SERVICE_ENV[service]
                 if os.getenv(key, "").strip()), DEFAULTS.get(service, ""))
    if not base:
        raise HTTPException(503, f"Real {service} upstream is not configured")
    return base.rstrip("/")


async def proxy_service(service: str, path: str, request: Request | None = None,
                        *, method: str = "GET", body: bytes | None = None) -> Response:
    base = service_base(service)
    headers = {}
    if request:
        for name in ("authorization", "content-type", "x-api-key", "x-tenant-id"):
            if value := request.headers.get(name):
                headers[name] = value
    try:
        async with httpx.AsyncClient(timeout=90 if "/export" in path else 5) as client:
            upstream = await client.request(
                method, base + path, headers=headers,
                params=list(request.query_params.multi_items()) if request else None,
                content=body if body is not None else
                    (await request.body() if request and method == "POST" else None),
            )
    except httpx.RequestError as exc:
        raise HTTPException(503, f"{service} upstream unavailable") from exc
    forwarded = {name: upstream.headers[name] for name in
                 ("content-type", "content-disposition", "cache-control")
                 if name in upstream.headers}
    return Response(upstream.content, status_code=upstream.status_code, headers=forwarded)


async def asi_snapshot() -> dict:
    async def probe(service: str):
        try:
            response = await proxy_service(service, "/health")
            return service, {"reachable": 200 <= response.status_code < 300,
                             "status_code": response.status_code}
        except HTTPException as exc:
            return service, {"reachable": False, "status_code": exc.status_code,
                             "error": exc.detail}
    services = dict(await asyncio.gather(*(probe(s) for s in ("alba", "albi", "jona"))))
    return {"status": "operational" if all(s["reachable"] for s in services.values())
            else "degraded", "services": services,
            "timestamp": datetime.now(timezone.utc).isoformat()}


@router.get("/api/asi/{service_name}/health")
async def asi_service_health(service_name: str):
    if service_name not in ("alba", "albi", "jona"):
        raise HTTPException(404, "Unknown ASI service")
    return await proxy_service(service_name, "/health")


@router.get("/api/proxy/marketplace-health")
async def marketplace_health(request: Request):
    return await proxy_service("marketplace", "/health", request)


@router.post("/api/query")
async def ocean_query(request: Request):
    return await proxy_service("ocean", "/api/v1/chat", request, method="POST")


@router.get("/api/zurich")
async def zurich_health(request: Request):
    return await proxy_service("ocean", "/api/v1/zurich/info", request)


@router.post("/api/zurich")
async def zurich_query(request: Request):
    return await proxy_service("ocean", "/api/v1/zurich", request, method="POST")


@router.get("/api/reporting/metrics")
async def reporting_metrics(request: Request):
    return await proxy_service("reporting", "/api/reporting/system-metrics", request)


@router.get("/api/reporting/alerts")
async def reporting_alerts(request: Request):
    return await proxy_service("reporting", "/api/reporting/alerts", request)


@router.get("/api/reporting/docker-stats")
async def reporting_docker_stats(request: Request):
    return await proxy_service("reporting", "/api/reporting/docker-stats", request)


@router.get("/api/reporting/export-pdf")
@router.post("/api/reporting/export-pdf")
async def reporting_pdf(request: Request):
    return await live_document_export(request, "pdf")


@router.get("/api/reporting/export-docx")
@router.post("/api/reporting/export-docx")
async def reporting_docx(request: Request):
    return await live_document_export(request, "docx")


async def live_document_export(request: Request, format: str) -> Response:
    upstream = await proxy_service("reporting", "/api/reporting/dashboard", request)
    if not 200 <= upstream.status_code < 300:
        return upstream
    try:
        data = json.loads(upstream.body)
    except (ValueError, UnicodeDecodeError) as exc:
        raise HTTPException(502, "Reporting dashboard returned invalid JSON") from exc
    output = io.BytesIO()
    if format == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        pdf = canvas.Canvas(output, pagesize=A4)
        width, height = A4
        y = height - 48
        pdf.setTitle("Clisonix live reporting")
        for line in ["Clisonix live reporting", *json.dumps(data, indent=2, ensure_ascii=True).splitlines()]:
            for wrapped in textwrap.wrap(line, width=95, replace_whitespace=False) or [""]:
                if y < 48:
                    pdf.showPage()
                    y = height - 48
                pdf.setFont("Courier", 9)
                pdf.drawString(40, y, wrapped)
                y -= 12
        pdf.save()
        media_type = "application/pdf"
    else:
        from docx import Document
        document = Document()
        document.add_heading("Clisonix live reporting", 0)
        document.add_paragraph(json.dumps(data, indent=2, ensure_ascii=False))
        document.save(output)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return Response(output.getvalue(), media_type=media_type, headers={
        "Content-Disposition": f'attachment; filename="clisonix-live-report.{format}"',
        "Cache-Control": "no-store",
    })


@router.get("/api/reporting/export")
@router.post("/api/reporting/export")
async def reporting_export(request: Request, format: str = Query("xlsx")):
    paths = {"xlsx": "export-excel", "pptx": "export-pptx",
             "pdf": "export-pdf", "docx": "export-docx"}
    if format not in paths:
        raise HTTPException(422, "Supported formats: xlsx, pptx, pdf, docx")
    if format in ("pdf", "docx"):
        return await live_document_export(request, format)
    return await proxy_service("reporting", "/api/reporting/" + paths[format],
                               request)


@router.get("/neural-symphony", status_code=410)
async def retired_synthetic_symphony():
    return JSONResponse({"error": "synthetic_eeg_endpoint_retired",
                         "message": "Use real EEG input through the EEG processing API.",
                         "input_endpoint": "/api/uploads/eeg/process"}, status_code=410)


def register_compatibility_routes(app: FastAPI) -> None:
    from reporting_api import get_unified_dashboard, get_errors, get_error_summary
    from unified_status_layer import unified_router
    from fastapi.routing import APIRoute

    # Reuse existing handlers, including their dependency-based authorization.
    for source, target in (("/", "/api"), ("/health", "/api/health"),
                           ("/status", "/api/status")):
        route = next(r for r in app.routes if isinstance(r, APIRoute)
                     and r.path == source and "GET" in r.methods)
        app.add_api_route(target, route.endpoint, methods=["GET"],
                          dependencies=route.dependencies, response_model=route.response_model,
                          response_class=route.response_class, tags=["runtime-compatibility"])
    for path, handler in (
        ("/api/proxy/reporting-dashboard", get_unified_dashboard),
        ("/api/proxy/reporting-errors", get_errors),
        ("/api/proxy/reporting-error-summary", get_error_summary),
    ):
        app.add_api_route(path, handler, methods=["GET"], tags=["reporting-compat"])
    app.include_router(unified_router)
    app.include_router(router)
    app.openapi_schema = None
